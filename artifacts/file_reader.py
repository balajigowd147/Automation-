from pathlib import Path

DEFAULT_MAX_CHARS = 5000

def read_txt(path, max_chars=DEFAULT_MAX_CHARS):

    path = Path(path)

    text = path.read_text(
        encoding="utf-8",
        errors="ignore",
    )
    if max_chars is None:
        return text

    return text[:max_chars]

def read_pdf(path, max_chars=DEFAULT_MAX_CHARS):

    from pypdf import PdfReader

    path = Path(path)

    reader = PdfReader(str(path))

    parts = []

    total_chars = 0

    for page in reader.pages:

        try:
            text = page.extract_text() or ""

        except Exception as error:

            print(
                f"[file_reader] Could not read PDF page: "
                f"{error}"
            )

            continue

        if max_chars is None:

            parts.append(text)

            continue

        remaining = max_chars - total_chars

        if remaining <= 0:
            break

        text = text[:remaining]

        parts.append(text)

        total_chars += len(text)

    result = "\n".join(parts)

    if max_chars is None:
        return result

    return result[:max_chars]

def read_docx(path, max_chars=DEFAULT_MAX_CHARS):

    from docx import Document

    path = Path(path)

    document = Document(str(path))

    parts = []

    total_chars = 0

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if not text:
            continue

        if max_chars is None:

            parts.append(text)

            continue

        remaining = max_chars - total_chars

        if remaining <= 0:
            break

        text = text[:remaining]

        parts.append(text)

        total_chars += len(text)

    result = "\n".join(parts)

    if max_chars is None:
        return result

    return result[:max_chars]

def read_pptx(path, max_chars=DEFAULT_MAX_CHARS):
    from pptx import Presentation

    path = Path(path)

    presentation = Presentation(str(path))

    parts = []

    total_chars = 0

    for slide in presentation.slides:

        for shape in slide.shapes:

            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if not text:
                continue

            if max_chars is None:

                parts.append(text)

                continue

            remaining = max_chars - total_chars

            if remaining <= 0:
                break

            text = text[:remaining]

            parts.append(text)

            total_chars += len(text)

        if max_chars is not None and total_chars >= max_chars:
            break

    result = "\n".join(parts)

    if max_chars is None:
        return result

    return result[:max_chars]

def read_xlsx(path, max_chars=DEFAULT_MAX_CHARS):

    from openpyxl import load_workbook

    path = Path(path)

    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )

    parts = []

    total_chars = 0

    try:

        for worksheet in workbook.worksheets:

            for row in worksheet.iter_rows(
                values_only=True
            ):

                values = []

                for value in row:

                    if value is None:
                        continue

                    values.append(
                        str(value)
                    )

                if not values:
                    continue

                text = " | ".join(values)

                if max_chars is None:

                    parts.append(text)

                    continue

                remaining = (
                    max_chars - total_chars
                )

                if remaining <= 0:
                    break

                text = text[:remaining]

                parts.append(text)

                total_chars += len(text)

            if max_chars is not None and total_chars >= max_chars:
                break

    finally:

        workbook.close()

    result = "\n".join(parts)

    if max_chars is None:
        return result

    return result[:max_chars]

def read_file(
    path,
    max_chars=DEFAULT_MAX_CHARS,
):

    path = Path(path)

    if not path.exists():

        raise FileNotFoundError(
            f"File does not exist: {path}"
        )

    if not path.is_file():

        raise ValueError(
            f"Path is not a file: {path}"
        )

    extension = path.suffix.lower()

    if extension == ".txt":

        return read_txt(
            path,
            max_chars,
        )

    if extension == ".pdf":

        return read_pdf(
            path,
            max_chars,
        )

    if extension == ".docx":

        return read_docx(
            path,
            max_chars,
        )

    if extension == ".pptx":

        return read_pptx(
            path,
            max_chars,
        )

    if extension == ".xlsx":

        return read_xlsx(
            path,
            max_chars,
        )

    raise ValueError(
        f"Unsupported file type: {extension}"
    )

def read_artifact(
    registry,
    artifact_id,
    max_chars=DEFAULT_MAX_CHARS,
):

    artifact = registry.get(
        artifact_id
    )

    if artifact is None:

        raise ValueError(
            f"Artifact not found: {artifact_id}"
        )

    if artifact["status"] != "active":

        raise ValueError(
            f"Artifact is not active: "
            f"{artifact_id}"
        )

    path = Path(
        artifact["path"]
    )

    return read_file(
        path,
        max_chars=max_chars,
    )

def update_artifact_snippet(
    registry,
    artifact_id,
    max_chars=1000,
):

    text = read_artifact(
        registry,
        artifact_id,
        max_chars=max_chars,
    )

    registry.update_content_snippet(
        artifact_id,
        text,
    )

    return text